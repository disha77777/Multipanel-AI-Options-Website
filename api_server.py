"""FastAPI backend that exposes the Multipanel Python tools over HTTP for the React UI."""

import io
import os
import re
import uuid
from pathlib import Path

import pandas as pd
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from pptx import Presentation as PptxDocument
from pydantic import BaseModel, Field

from excel import add_workbook_chart, chart as create_excel_chart_preview, export_research, parse_research_response, research_query
from datacleaner import apply_operations, clean_pipeline, interpret_request, summarize
from presentation import powerpoint, presentation as build_pptx, refine_presentation
from reader import process_image_ocr
from research_assistant import research_verified
from research_planner import generate_research_plan, refine_research_plan
from ai_provider import configured_providers

BASE_DIR = Path(__file__).resolve().parent
MAX_IMAGE_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_DATA_UPLOAD_BYTES = 20 * 1024 * 1024


def load_key_env():
    """Load local key.env / .env values without requiring an additional package."""
    for filename in ("key.env", ".env"):
        env_file = BASE_DIR / filename
        if not env_file.exists():
            continue
        for line in env_file.read_text().splitlines():
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            os.environ[key.strip()] = value.strip().strip('"').strip("'")


load_key_env()

app = FastAPI(title="Multipanel API")

allowed_origins = [
    origin.strip()
    for origin in os.getenv("ALLOWED_ORIGINS", "http://localhost:8443,http://127.0.0.1:8443").split(",")
    if origin.strip()
]
app.add_middleware(
    CORSMiddleware,
    allow_origins=allowed_origins,
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

# In-memory session store for the data cleaner (single-process demo state).
_cleaner_sessions: dict[str, pd.DataFrame] = {}
_excel_sessions: dict[str, dict] = {}


def _remember_excel(dataframe, *, summary="", sources=None, query="", uploaded=False):
    session_id = uuid.uuid4().hex
    if len(_excel_sessions) >= 12:
        _excel_sessions.pop(next(iter(_excel_sessions)))
    _excel_sessions[session_id] = {
        "dataframe": dataframe.copy(), "summary": summary, "sources": sources or [],
        "query": query, "uploaded": uploaded,
    }
    return session_id


def friendly_error(error: Exception) -> str:
    message = str(error)
    if "429" in message or "RESOURCE_EXHAUSTED" in message:
        return (
            "Gemini quota reached (429). The configured Google project or selected model "
            "has reached a request, token, daily, or billing limit. Check the correct "
            "project's active limits and usage at https://aistudio.google.com/usage, and "
            "confirm billing and the model's tier. A short wait can resolve a temporary "
            "rate limit; daily or billing limits need a reset or account change. "
            "The error did not include enough quota details to identify which limit was hit."
        )
    return message


def df_preview(df: pd.DataFrame, limit: int = 200):
    return {
        "columns": list(df.columns.astype(str)),
        "rows": df.head(limit).astype(object).where(pd.notna(df.head(limit)), None).to_dict(orient="records"),
        "total_rows": len(df),
    }


def parse_pptx(file_path: Path) -> dict:
    """Read an existing .pptx file into the {title, slides:[{title, bullets}]} shape."""
    deck = PptxDocument(str(file_path))
    slides = []
    deck_title = None
    for index, slide in enumerate(deck.slides):
        title_shape = slide.shapes.title
        title_text = title_shape.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        bullets = []
        for shape in slide.shapes:
            if shape is title_shape or not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                text = paragraph.text.strip()
                if text:
                    bullets.append(text)
        if index == 0 and title_text:
            deck_title = title_text
        slides.append({"title": title_text or f"Slide {index + 1}", "bullets": bullets})
    if not slides:
        raise ValueError("This presentation has no slides.")
    return {"title": deck_title or "Uploaded presentation", "slides": slides}


# ─── Presentation Studio ──────────────────────────────────────────────────────

class PresentationRequest(BaseModel):
    topic: str = Field(min_length=1, max_length=12000)
    slide_count: int = Field(default=6, ge=1, le=30)
    audience: str = Field(default="general audience", max_length=300)
    tone: str = Field(default="clear and persuasive", max_length=300)
    provider: str = Field(default="grok", max_length=20)
    style: str = Field(default="Editorial", max_length=40)
    include_web_images: bool = False


class PresentationRefineRequest(BaseModel):
    presentation: dict
    revision: str = Field(min_length=1, max_length=12000)
    provider: str = Field(default="grok", max_length=20)
    style: str = Field(default="Editorial", max_length=40)
    include_web_images: bool = False


@app.post("/api/presentation")
def create_presentation(payload: PresentationRequest):
    if not payload.topic.strip():
        raise HTTPException(400, "Enter a topic or thesis statement first.")
    try:
        return powerpoint(
            user_query=payload.topic.strip(),
            slide_count=payload.slide_count,
            audience=payload.audience.strip() or "general audience",
            tone=payload.tone.strip() or "clear and persuasive",
            provider=payload.provider,
            style=payload.style,
            include_web_images=payload.include_web_images,
        )
    except Exception as error:
        raise HTTPException(502, friendly_error(error)) from error


@app.post("/api/presentation/refine")
def refine_presentation_endpoint(payload: PresentationRefineRequest):
    try:
        return refine_presentation(
            payload.presentation, payload.revision, provider=payload.provider,
            style=payload.style, include_web_images=payload.include_web_images,
        )
    except Exception as error:
        raise HTTPException(502, friendly_error(error)) from error


@app.post("/api/presentation/upload")
async def upload_presentation_file(file: UploadFile = File(...)):
    suffix = Path(file.filename or "deck.pptx").suffix or ".pptx"
    if suffix.lower() != ".pptx":
        raise HTTPException(400, "Only .pptx files can be uploaded here.")
    temp_path = BASE_DIR / f"_upload_{uuid.uuid4().hex}{suffix}"
    try:
        deck_bytes = await file.read(MAX_DATA_UPLOAD_BYTES + 1)
        if len(deck_bytes) > MAX_DATA_UPLOAD_BYTES:
            raise HTTPException(413, "Presentation is too large. Upload a .pptx file under 20 MB.")
        temp_path.write_bytes(deck_bytes)
        return parse_pptx(temp_path)
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(400, f"Could not read this presentation: {friendly_error(error)}") from error
    finally:
        temp_path.unlink(missing_ok=True)



@app.post("/api/presentation/export")
def export_presentation(payload: dict):
    try:
        output_path = build_pptx(payload, ppt_file=str(BASE_DIR / "report.pptx"))
        return FileResponse(
            output_path,
            filename="presentation.pptx",
            media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )
    except Exception as error:
        raise HTTPException(502, friendly_error(error)) from error


# ─── Excel Research ────────────────────────────────────────────────────────────

class ExcelRequest(BaseModel):
    query: str = Field(min_length=1, max_length=12000)
    provider: str = Field(default="grok", max_length=20)


class ExcelChartRequest(BaseModel):
    session_id: str | None = Field(default=None, max_length=80)
    columns: list[str] | None = Field(default=None, min_length=1, max_length=100)
    rows: list[dict] | None = Field(default=None, min_length=1, max_length=200)
    total_rows: int = Field(default=0, ge=0, le=1000000)
    summary: str = Field(default="", max_length=12000)
    sources: list[str] = Field(default_factory=list, max_length=8)
    query: str = Field(default="", max_length=12000)
    chart_request: str = Field(min_length=1, max_length=1000)


@app.post("/api/excel")
def create_excel(payload: ExcelRequest):
    if not payload.query.strip():
        raise HTTPException(400, "Enter a research question first.")
    try:
        response_text = research_query(payload.query.strip(), provider=payload.provider)
        summary, sources, dataframe = parse_research_response(response_text)
        output_path = BASE_DIR / "ai_table.xlsx"
        export_research(dataframe, excel_file=output_path, summary=summary, sources=sources, query=payload.query.strip())
        chart_path = create_excel_chart_preview(dataframe, output_file=str(BASE_DIR / "ai_chart.png"), query=payload.query.strip())
        session_id = _remember_excel(dataframe, summary=summary, sources=sources, query=payload.query.strip())
        return {"summary": summary, "sources": sources, "query": payload.query.strip(), "session_id": session_id,
                "chart_available": bool(chart_path), **df_preview(dataframe)}
    except Exception as error:
        raise HTTPException(502, friendly_error(error)) from error


@app.post("/api/excel/chart")
def refine_excel_chart(payload: ExcelChartRequest):
    """Refresh the current workbook chart from its visible data without another AI call."""
    try:
        session = _excel_sessions.get(payload.session_id or "")
        if session:
            resolved_session_id = payload.session_id
            dataframe = session["dataframe"]
            summary, sources, query = session["summary"], session["sources"], session["query"]
        elif payload.columns and payload.rows:
            # Handle a workbook already open in the browser from before a page refresh.
            dataframe = pd.DataFrame(payload.rows, columns=payload.columns)
            if payload.total_rows > len(dataframe):
                raise HTTPException(422, "This workbook has more rows than the current preview. Reload it before charting to include every row.")
            summary, sources, query = payload.summary, payload.sources, payload.query
            resolved_session_id = _remember_excel(dataframe, summary=summary, sources=sources, query=query)
            session = _excel_sessions[resolved_session_id]
        else:
            raise HTTPException(404, "This workbook session has expired. Reload or upload the workbook, then try again.")

        chart_request = payload.chart_request.strip()
        lower_request = chart_request.lower()
        dimension_aliases = {
            "region": ("region", "regions", "regional"),
            "country": ("country", "countries"),
            "company": ("company", "companies", "organization", "entity"),
            "period": ("period", "year", "date", "time"),
            "category": ("category", "segment", "market"),
        }
        columns_lower = {str(column).strip().lower() for column in dataframe.columns}
        for dimension, aliases in dimension_aliases.items():
            if any(re.search(rf"\b{re.escape(alias)}\b", lower_request) for alias in aliases):
                if not (dimension in columns_lower or any(alias in columns_lower for alias in aliases)):
                    available = ", ".join(str(column) for column in dataframe.columns)
                    raise HTTPException(422, f"This workbook has no {dimension} column. Available fields: {available}.")
                break

        chart_query = f"{query}\nChart request: {chart_request}".strip()
        output_path = BASE_DIR / "ai_table.xlsx"
        chart_path = BASE_DIR / "ai_chart.png"
        preview_path = create_excel_chart_preview(dataframe, output_file=str(chart_path), query=chart_query)
        if not preview_path:
            raise HTTPException(422, "This data has no single comparable numeric measure to chart.")
        if session.get("uploaded"):
            if not add_workbook_chart(output_path, dataframe, query=chart_query, identifier=resolved_session_id or "local"):
                raise HTTPException(422, "This data has no single comparable numeric measure to chart.")
        else:
            export_research(dataframe, excel_file=output_path, summary=summary,
                            sources=sources, query=query, chart_query=chart_query)
        return {"summary": summary, "sources": sources, "query": query, "session_id": resolved_session_id,
                "chart_available": True, **df_preview(dataframe)}
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(422, f"Could not update this chart: {friendly_error(error)}") from error


@app.get("/api/excel/download")
def download_excel():
    output_path = BASE_DIR / "ai_table.xlsx"
    if not output_path.exists():
        raise HTTPException(404, "No workbook has been generated yet.")
    return FileResponse(output_path, filename="ai_table.xlsx")


@app.get("/api/excel/chart")
def download_excel_chart():
    chart_path = BASE_DIR / "ai_chart.png"
    if not chart_path.exists():
        raise HTTPException(404, "No chart is available for this workbook.")
    return FileResponse(chart_path, media_type="image/png", filename="ai_chart.png", headers={"Cache-Control": "no-store"})


@app.post("/api/excel/upload")
async def upload_excel_file(file: UploadFile = File(...)):
    name = (file.filename or "").lower()
    try:
        if not name.endswith((".csv", ".xlsx")):
            raise HTTPException(400, "Use a .xlsx workbook or .csv file. Save legacy .xls files as .xlsx first.")
        raw_bytes = await file.read(MAX_DATA_UPLOAD_BYTES + 1)
        if len(raw_bytes) > MAX_DATA_UPLOAD_BYTES:
            raise HTTPException(413, "Data file is too large. Upload a file under 20 MB.")
        if name.endswith(".csv"):
            dataframe = pd.read_csv(io.BytesIO(raw_bytes))
        else:
            sheets = pd.read_excel(io.BytesIO(raw_bytes), sheet_name=None)
            dataframe = sheets.get("Research Data", next(iter(sheets.values())))

        output_path = BASE_DIR / "ai_table.xlsx"
        if name.endswith(".xlsx"):
            output_path.write_bytes(raw_bytes)
        else:
            dataframe.to_excel(output_path, index=False)
        preview_path = create_excel_chart_preview(dataframe, output_file=str(BASE_DIR / "ai_chart.png"), query="uploaded spreadsheet")
        session_id = _remember_excel(dataframe, uploaded=True)
        if preview_path:
            add_workbook_chart(output_path, dataframe, query="uploaded spreadsheet", identifier=session_id)
        return {"summary": "Uploaded data. Review the table and chart before using the workbook.", "sources": [],
                "session_id": session_id, "chart_available": bool(preview_path), **df_preview(dataframe)}
    except Exception as error:
        if isinstance(error, HTTPException):
            raise
        raise HTTPException(400, f"Could not read this file: {friendly_error(error)}") from error


# ─── Verified Research ─────────────────────────────────────────────────────────

class ResearchRequest(BaseModel):
    question: str = Field(min_length=1, max_length=1500)
    date_range: str = Field(default="latest available", max_length=300)
    source_preference: str = Field(default="credible primary sources", max_length=500)
    provider: str = Field(default="grok", max_length=20)


@app.post("/api/research")
def create_verified_research(payload: ResearchRequest):
    if not payload.question.strip():
        raise HTTPException(400, "Enter a research question first.")
    try:
        return research_verified(
            payload.question.strip(),
            date_range=payload.date_range.strip() or "latest available",
            source_preference=payload.source_preference.strip() or "credible primary sources",
            provider=payload.provider,
        )
    except Exception as error:
        raise HTTPException(502, friendly_error(error)) from error


# ─── Research Planner ──────────────────────────────────────────────────────────

class PlannerRequest(BaseModel):
    thesis_statement: str = Field(min_length=1, max_length=12000)
    field_of_study: str = Field(default="Research", max_length=300)
    research_type: str = Field(default="Thesis/Dissertation", max_length=300)
    duration_months: int = Field(default=12, ge=1, le=120)
    provider: str = Field(default="grok", max_length=20)


class PlannerRefineRequest(BaseModel):
    plan: dict
    revision: str = Field(min_length=1, max_length=12000)
    provider: str = Field(default="grok", max_length=20)


@app.post("/api/planner")
def create_plan(payload: PlannerRequest):
    if not payload.thesis_statement.strip():
        raise HTTPException(400, "Enter a thesis statement or research question first.")
    try:
        return generate_research_plan(
            thesis_statement=payload.thesis_statement.strip(),
            field_of_study=payload.field_of_study.strip() or "Research",
            research_type=payload.research_type,
            duration_months=payload.duration_months,
            provider=payload.provider,
        )
    except Exception as error:
        raise HTTPException(502, friendly_error(error)) from error


@app.post("/api/planner/refine")
def refine_plan(payload: PlannerRefineRequest):
    try:
        return refine_research_plan(payload.plan, payload.revision, provider=payload.provider)
    except Exception as error:
        raise HTTPException(502, friendly_error(error)) from error


# ─── Data Cleaner ───────────────────────────────────────────────────────────────

class CleanerChatRequest(BaseModel):
    session_id: str = Field(min_length=1, max_length=80)
    message: str = Field(min_length=1, max_length=4000)
    provider: str = Field(default="grok", max_length=20)


@app.post("/api/cleaner/upload")
async def upload_cleaner_file(file: UploadFile = File(...)):
    try:
        name = (file.filename or "").lower()
        if not name.endswith((".csv", ".xlsx")):
            raise HTTPException(400, "Use a .xlsx workbook or .csv file. Save legacy .xls files as .xlsx first.")
        raw_bytes = await file.read(MAX_DATA_UPLOAD_BYTES + 1)
        if len(raw_bytes) > MAX_DATA_UPLOAD_BYTES:
            raise HTTPException(413, "Data file is too large. Upload a file under 20 MB.")
        if name.endswith(".csv"):
            raw_df = pd.read_csv(io.BytesIO(raw_bytes))
        else:
            raw_df = pd.read_excel(io.BytesIO(raw_bytes))

        cleaned_df = clean_pipeline(raw_df)
        session_id = str(uuid.uuid4())
        if len(_cleaner_sessions) >= 12:
            _cleaner_sessions.pop(next(iter(_cleaner_sessions)))
        _cleaner_sessions[session_id] = cleaned_df

        stats = summarize(cleaned_df)
        return {
            "session_id": session_id,
            "source_rows": len(raw_df),
            **df_preview(cleaned_df),
            "missing_by_column": stats["missing_by_column"],
            "dtypes": stats["dtypes"],
        }
    except Exception as error:
        if isinstance(error, HTTPException):
            raise
        raise HTTPException(400, f"Could not read this file: {friendly_error(error)}") from error


@app.post("/api/cleaner/chat")
def cleaner_chat(payload: CleanerChatRequest):
    df = _cleaner_sessions.get(payload.session_id)
    if df is None:
        raise HTTPException(404, "Upload a file first to start a cleaning session.")
    try:
        operations = interpret_request(df, payload.message, provider=payload.provider)
        if not operations:
            stats = summarize(df)
            return {
                "log": ["I couldn't turn that into a supported cleaning step. Try naming the column and the exact action."],
                **df_preview(df),
                "missing_by_column": stats["missing_by_column"],
                "dtypes": stats["dtypes"],
            }

        updated_df, log = apply_operations(df, operations)
        _cleaner_sessions[payload.session_id] = updated_df
        stats = summarize(updated_df)
        return {
            "log": log,
            **df_preview(updated_df),
            "missing_by_column": stats["missing_by_column"],
            "dtypes": stats["dtypes"],
        }
    except Exception as error:
        raise HTTPException(502, f"Could not apply that request: {friendly_error(error)}") from error


@app.get("/api/cleaner/download/{session_id}")
def download_cleaned(session_id: str):
    df = _cleaner_sessions.get(session_id)
    if df is None:
        raise HTTPException(404, "No cleaning session found.")
    csv_bytes = df.to_csv(index=False).encode("utf-8")
    return FileResponse(
        path=_write_temp_csv(csv_bytes),
        filename="cleaned_data.csv",
        media_type="text/csv",
    )


def _write_temp_csv(csv_bytes: bytes) -> Path:
    temp_path = BASE_DIR / f"_cleaned_{uuid.uuid4().hex}.csv"
    temp_path.write_bytes(csv_bytes)
    return temp_path


# ─── Document Reader ────────────────────────────────────────────────────────────

@app.post("/api/extractor/upload")
async def upload_extractor_image(file: UploadFile = File(...), provider: str = Form(default="grok")):
    suffix = (Path(file.filename or "image.jpg").suffix or ".jpg").lower()
    if suffix not in {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif", ".tif", ".tiff"}:
        raise HTTPException(400, "Use a JPG, PNG, WebP, BMP, GIF, or TIFF image.")
    temp_path = BASE_DIR / f"_upload_{uuid.uuid4().hex}{suffix}"
    try:
        image_bytes = await file.read(MAX_IMAGE_UPLOAD_BYTES + 1)
        if len(image_bytes) > MAX_IMAGE_UPLOAD_BYTES:
            raise HTTPException(413, "Image is too large. Upload an image under 20 MB.")
        temp_path.write_bytes(image_bytes)
        output_path = BASE_DIR / "ocr_output.xlsx"
        markdown_text, df, saved_path = process_image_ocr(str(temp_path), str(output_path), provider=provider)

        result = {"markdown": markdown_text, "columns": [], "rows": [], "total_rows": 0}
        if df is not None:
            df.to_csv(BASE_DIR / "ocr_output.csv", index=False)
            result.update(df_preview(df))
        return result
    except HTTPException:
        raise
    except ValueError as error:
        raise HTTPException(400, str(error)) from error
    except Exception as error:
        raise HTTPException(502, friendly_error(error)) from error
    finally:
        temp_path.unlink(missing_ok=True)


@app.get("/api/extractor/download/{fmt}")
def download_extraction(fmt: str):
    if fmt not in ("csv", "xlsx"):
        raise HTTPException(400, "Format must be csv or xlsx.")
    output_path = BASE_DIR / f"ocr_output.{fmt}"
    if not output_path.exists():
        raise HTTPException(404, "No extraction has been generated yet.")
    return FileResponse(output_path, filename=f"ocr_output.{fmt}")


@app.get("/api/health")
def health():
    providers = configured_providers()
    return {"status": "ok", "ai_configured": any(providers.values()), "providers": providers}
